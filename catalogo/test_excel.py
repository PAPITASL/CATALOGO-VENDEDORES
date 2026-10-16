from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .models import Marca, ModeloVehiculo, Producto


class ExcelTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username='excel-admin', is_staff=True)
        self.client.force_login(self.user)
        brand = Marca.objects.create(nombre='Chevrolet')
        model = ModeloVehiculo.objects.create(marca=brand, nombre='Silverado')
        Producto.objects.bulk_create([
            Producto(modelo=model, nombre_pieza='Persiana cromada', sku=f'{n:04}',
                     anio_inicio=1988, anio_fin=1998, precio_minimo=Decimal('100.25'),
                     precio_venta=Decimal('120.50'), costo=None, observaciones='=1+1',
                     activo=n != 0, imagen_principal='photo.jpg')
            for n in range(30)
        ])
        self.url = reverse('catalogo:download_excel')

    def workbook(self, query=''):
        response = self.client.get(self.url + query)
        self.assertEqual(response.status_code, 200)
        self.assertIn('.xlsx', response['Content-Disposition'])
        self.assertEqual(response['Cache-Control'], 'no-store')
        book = load_workbook(BytesIO(b''.join(response.streaming_content)))
        self.addCleanup(book.close)
        return book.active

    def test_all_pages_types_and_literal_text(self):
        sheet = self.workbook('?page=2')
        self.assertEqual(sheet.max_row, 31)
        self.assertEqual(sheet['A2'].value, 'Chevrolet')
        self.assertEqual(sheet['B2'].value, 'Silverado')
        self.assertEqual(sheet['C2'].value, 1988)
        self.assertEqual(sheet['D2'].value, 1998)
        self.assertEqual(sheet['F2'].value, '0000')
        self.assertIsNone(sheet['G2'].value)
        self.assertIsNone(sheet['H2'].value)
        self.assertEqual(sheet['I2'].value, 100.25)
        self.assertEqual(sheet['J2'].value, 120.5)
        self.assertEqual(sheet['J2'].data_type, 'n')
        self.assertEqual(sheet['K2'].value, '=1+1')
        self.assertEqual(sheet['K2'].data_type, 's')
        self.assertEqual(sheet['L2'].value, 'Inactivo')
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:L31')

    def test_filters_and_empty_export(self):
        self.assertEqual(self.workbook('?activo=0&anio=1990&q=Persiana').max_row, 2)
        self.assertEqual(self.workbook('?q=inexistente').max_row, 1)

    def test_selection_limits_excel_and_catalog_including_inactive(self):
        from .views import catalog_download_products
        product = Producto.objects.get(sku='0000')
        sheet = self.workbook(f'?selected={product.pk}')
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['F2'].value, '0000')
        self.assertEqual(list(catalog_download_products({'selected': str(product.pk)})), [product])
        self.assertEqual(self.workbook('?selected=invalid').max_row, 1)
        self.assertEqual(self.workbook(f'?selected={product.pk}&activo=1').max_row, 1)

    def test_access_and_button(self):
        self.assertContains(self.client.get(reverse('catalogo:producto_list')), 'Descargar Excel (30)')
        self.assertEqual(self.client.post(self.url).status_code, 405)
        self.user.is_staff = False
        self.user.save()
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
