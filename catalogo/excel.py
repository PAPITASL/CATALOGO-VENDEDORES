from tempfile import SpooledTemporaryFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill


def build_products_excel(products):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('Productos')
    sheet.freeze_panes = 'A2'
    headers = ['Marca', 'Modelo', 'Año inicial', 'Año final', 'Accesorio', 'SKU',
               'Categoría', 'Costo', 'Precio mínimo', 'Precio de venta',
               'Observaciones', 'Estado']
    widths = [20, 28, 14, 14, 45, 20, 24, 20, 20, 20, 65, 16]
    for column, width in zip('ABCDEFGHIJKL', widths):
        sheet.column_dimensions[column].width = width
    heading = []
    for label in headers:
        cell = WriteOnlyCell(sheet, value=label)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='23445C')
        heading.append(cell)
    sheet.append(heading)
    count = 1
    for product in products.iterator(chunk_size=200):
        values = [product.modelo.marca.nombre, product.modelo.nombre,
                  product.anio_inicio, product.anio_fin, product.nombre_pieza,
                  product.sku, product.categoria.nombre if product.categoria else None,
                  product.costo, product.precio_minimo, product.precio_venta,
                  product.observaciones, 'Activo' if product.activo else 'Inactivo']
        row = []
        for index, value in enumerate(values):
            if isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub('', value)
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                # Product text, including leading '=' and zero-prefixed SKUs,
                # must stay literal, never become a spreadsheet formula.
                cell.data_type = 's'
            if index in (7, 8, 9):
                cell.number_format = '#,##0.00'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            row.append(cell)
        sheet.append(row)
        count += 1
    sheet.auto_filter.ref = f'A1:L{count}'
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode='w+b')
    try:
        workbook.save(output)
        output.seek(0)
        return output
    except BaseException:
        output.close()
        raise
    finally:
        workbook.close()
